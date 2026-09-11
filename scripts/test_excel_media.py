"""Real HTTP checks using the disposable QA database and temporary image storage."""
from io import BytesIO
from tempfile import TemporaryDirectory
from pathlib import Path
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from qa_backend import app, ADMIN_EMAIL, CLIENT_EMAIL, PASSWORD
from src.infrastructure.config.settings import settings
from src.shared.bulk.registry import RESOURCES

client = TestClient(app)
def login(email):
    result = client.post('/api/v1/auth/login', json={'email': email, 'password': PASSWORD})
    assert result.status_code == 200, result.text
    return {'Authorization': 'Bearer ' + result.json()['data']['access_token']}
admin, customer = login(ADMIN_EMAIL), login(CLIENT_EMAIL)
for resource in RESOURCES:
    for action in ('template', 'export'):
        result = client.get(f'/api/v1/bulk/{resource}/{action}', headers=admin)
        assert result.status_code == 200, (resource, action, result.text)
        book = load_workbook(BytesIO(result.content))
        assert 'Datos' in book.sheetnames
        book.close()
assert client.get('/api/v1/bulk/sizes/template', headers=customer).status_code == 403
template = client.get('/api/v1/bulk/sizes/template', headers=admin)
book = load_workbook(BytesIO(template.content))
book['Datos'].append(['QA-EXCEL', 'Talla Excel', 50])
stream = BytesIO()
book.save(stream)
book.close()
payload = stream.getvalue()
files = {'file': ('sizes.xlsx', payload, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
preview = client.post('/api/v1/bulk/sizes/preview', headers=admin, files=files, data={'mode':'create'})
assert preview.status_code == 200, preview.text
report = preview.json()['data']
assert report['valid'] == 1 and report['imported'] == 0, report
def sizes():
    return client.get('/api/v1/catalog/admin/sizes', headers=admin).json()['data']
assert not any(row['code'] == 'QA-EXCEL' for row in sizes()), 'Preview persisted a record'
result = client.post('/api/v1/bulk/sizes/import', headers=admin, files=files,
                     data={'mode':'create', 'confirm':'true', 'preview_digest':report['digest']})
assert result.status_code == 200 and result.json()['data']['imported'] == 1, result.text
assert any(row['code'] == 'QA-EXCEL' for row in sizes())
duplicate = client.post('/api/v1/bulk/sizes/preview', headers=admin, files=files, data={'mode':'create'})
assert duplicate.json()['data']['errors'], duplicate.text

# A valid row followed by a duplicate must not partially persist.
book = load_workbook(BytesIO(template.content))
book['Datos'].append(['QA-ATOMIC', 'No debe guardarse', 51])
book['Datos'].append(['QA-EXCEL', 'Duplicada', 52])
stream = BytesIO()
book.save(stream)
book.close()
mixed_files = {'file': ('mixed.xlsx', stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
mixed_preview = client.post('/api/v1/bulk/sizes/preview', headers=admin, files=mixed_files, data={'mode':'create'})
mixed_report = mixed_preview.json()['data']
assert mixed_report['errors'] and mixed_report['valid'] == 1, mixed_report
mixed_import = client.post('/api/v1/bulk/sizes/import', headers=admin, files=mixed_files,
                          data={'mode':'create','confirm':'true','preview_digest':mixed_report['digest']})
assert mixed_import.status_code == 200 and mixed_import.json()['data']['imported'] == 0, mixed_import.text
assert not any(row['code'] == 'QA-ATOMIC' for row in sizes()), 'Partial import persisted a row'
changed = client.post('/api/v1/bulk/sizes/import', headers=admin, files=mixed_files,
                      data={'mode':'create','confirm':'true','preview_digest':report['digest']})
assert 400 <= changed.status_code < 500, changed.text
book = load_workbook(BytesIO(template.content))
book['Datos'].append(['QA-FORMULA', '=1+1', 53])
stream = BytesIO()
book.save(stream)
book.close()
formula = client.post('/api/v1/bulk/sizes/preview', headers=admin,
                      files={'file': ('formula.xlsx', stream.getvalue())}, data={'mode':'create'})
assert 400 <= formula.status_code < 500, formula.text

previous = settings.media_storage_dir
try:
    with TemporaryDirectory() as directory:
        settings.media_storage_dir = directory
        stream = BytesIO()
        Image.new('RGB', (20, 30), 'red').save(stream, 'PNG')
        upload = {'file': ('photo.png', stream.getvalue(), 'image/png')}
        assert client.post('/api/v1/media/images', headers=customer, files=upload).status_code == 403
        result = client.post('/api/v1/media/images', headers=admin, files=upload)
        assert result.status_code == 201, result.text
        data = result.json()['data']
        assert (data['width'], data['height']) == (20, 30)
        assert (Path(directory) / data['name']).is_file()
        public = client.get('/api/v1/media/files/' + data['name'])
        assert public.status_code == 200 and public.headers['content-type'] == 'image/webp'
        invalid = client.post('/api/v1/media/images', headers=admin,
                              files={'file': ('fake.png', b'not an image', 'image/png')})
        assert 400 <= invalid.status_code < 500, invalid.text
finally:
    settings.media_storage_dir = previous
print('Excel: templates/exports, permissions, preview rollback, import and duplicates passed. Images: upload, public retrieval, permissions and invalid content passed.')
